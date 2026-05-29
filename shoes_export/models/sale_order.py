# Copyright Serincloud SL - 2025
# License AGPL-3 - See http://www.gnu.org/licenses/agpl-3.0.html
import base64
import io
import unicodedata
from datetime import date

from odoo import _, models
from odoo.exceptions import UserError

try:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as XlImage
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    Workbook = None
    XlImage = None

_FILL_HEADER = PatternFill("solid", fgColor="1F4E79") if Workbook else None
_FILL_ALT = PatternFill("solid", fgColor="DEEAF1") if Workbook else None
_FONT_WHITE_BOLD = Font(bold=True, color="FFFFFF") if Workbook else None
_ALIGN_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True) if Workbook else None
_ALIGN_LEFT   = Alignment(horizontal="left",   vertical="center") if Workbook else None
_ALIGN_RIGHT  = Alignment(horizontal="right",  vertical="center") if Workbook else None

# (key, label, col_width, fmt)  fmt: text | integer | number | currency | date
_EAN_COLS = [
    ("num_pedido",  "Núm pedido",          18, "text"),
    ("marca",       "Marca",               16, "text"),
    ("temporada",   "Temporada",           18, "text"),
    ("nombre_web",  "NombreWeb",           25, "text"),
    ("modelo",      "Modelo",              30, "text"),
    ("color",       "Color",               14, "text"),
    ("talla",       "Talla",               10, "text"),
    ("quantity",    "quantity",            12, "integer"),
    ("tip",         "Tip",                 16, "text"),
    ("barcode",     "Código de barras",    20, "text"),
    ("desc_corta",  "Desc corta",          30, "text"),
    ("mat_ext",     "Material Exterior",   30, "text"),
    ("plantilla",   "Plantilla",           30, "text"),
    ("forro",       "Forro",               18, "text"),
    ("punta",       "Punta",               14, "text"),
    ("tacon",       "Tacón",               14, "text"),
    ("alt_tacon",   "Altura Tacón",        12, "number"),
    ("coste",       "Coste",               14, "currency"),
    ("moneda",      "Moneda",              10, "text"),
    ("origen",      "Origen",              10, "text"),
    ("arancel",     "Partida Arancelaria", 22, "text"),
]

_SURTIDOS_COLS = [
    ("imagen",      "Imagen",              12, "image"),
    ("navima_po",   "NAVIMA PO NR.",       18, "text"),
    ("temporada",   "temporada",           18, "text"),
    ("brand",       "brand",               16, "text"),
    ("modelo",      "modelo",              30, "text"),
    ("surtido",     "surtido",             16, "text"),
    ("style_code",  "STYLE CODE",          35, "text"),
    ("color",       "COLOR",               14, "text"),
    ("categoria",   "categoría",           18, "text"),
    ("price",       "price",               14, "currency"),
    ("moneda",      "moneda",              10, "text"),
    ("tallas",      "Tallas",              40, "text"),
    ("par_nota",    "par_nota",            12, "integer"),
    ("bultos",      "bultos",              10, "integer"),
    ("pares_bul",   "pares_bul",           12, "integer"),
    ("portes",      "portes",              14, "text"),
    ("su_ref",      "su_referencia",       25, "text"),
]


class SaleOrder(models.Model):
    _inherit = "sale.order"

    def action_export_shoes_excel(self):
        if not Workbook:
            raise UserError(_("La librería openpyxl no está disponible en este entorno."))

        commercial_partners = self.mapped("partner_id.commercial_partner_id")
        if len(commercial_partners) > 1:
            names = ", ".join(commercial_partners.mapped("name"))
            raise UserError(
                _("Los pedidos seleccionados pertenecen a distintos clientes comerciales (%s). "
                  "Selecciona solo pedidos del mismo cliente para exportar juntos.") % names
            )

        today = date.today().strftime("%Y%m%d")
        doc_name = self.name if len(self) == 1 else "pedidos_shoes"
        raw = f"{self.env.company.name}_{today}_{doc_name}"
        normalized = unicodedata.normalize("NFKD", raw)
        filename = (
            "".join(c for c in normalized if not unicodedata.combining(c))
            .replace(" ", "_") + ".xlsx"
        )

        wb = Workbook()
        wb.remove(wb.active)

        lines = self.env["sale.order.line"]
        for order in self:
            lines |= order.order_line.filtered(lambda l: not l.display_type)

        self._shoes_write_ean(wb, lines)
        self._shoes_write_surtidos(wb, lines)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        attachment = self.env["ir.attachment"].sudo().create({
            "name": filename,
            "type": "binary",
            "datas": base64.b64encode(buf.read()),
            "mimetype": (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            "res_model": self._name,
            "res_id": self[0].id,
        })
        return {
            "type": "ir.actions.act_url",
            "url": f"/web/content/{attachment.id}?download=true",
            "target": "self",
        }

    # -------------------------------------------------------------------------
    # Hoja EAN
    # -------------------------------------------------------------------------

    def _shoes_write_ean(self, wb, lines):
        ws = wb.create_sheet("EAN")
        ws.freeze_panes = "A2"
        self._shoes_header(ws, _EAN_COLS)

        for row_idx, row in enumerate(self._shoes_ean_rows(lines), start=2):
            fill = _FILL_ALT if row_idx % 2 == 0 else None
            for col, (key, _, _, fmt) in enumerate(_EAN_COLS, start=1):
                val = self._shoes_fmt(row.get(key, ""), fmt)
                cell = ws.cell(row=row_idx, column=col, value=val)
                if fill:
                    cell.fill = fill
                cell.number_format = self._shoes_numfmt(fmt, val)
                cell.alignment = _ALIGN_RIGHT if fmt in ("currency", "number", "integer") else _ALIGN_LEFT

    def _shoes_ean_rows(self, lines):
        # Agrega por (order_id, pair_product_id); conserva orden de aparición.
        agg = {}
        seq = []

        for line in lines:
            product = line.product_id
            order = line.order_id

            pairs = []  # (pair_product, qty)
            if product.is_assortment:
                bom = product.bom_ids[:1]
                if bom:
                    for bl in bom.bom_line_ids:
                        if bl.product_id.is_pair:
                            pairs.append((bl.product_id, bl.product_qty * line.product_uom_qty))
            elif product.is_pair:
                pairs.append((product, line.product_uom_qty))

            for pair, qty in pairs:
                key = (order.id, pair.id)
                if key not in agg:
                    seq.append(key)
                    task = self._shoes_task(pair)
                    # trade_name: del par; si vacío, del surtido padre
                    trade_name = (
                        getattr(pair, "trade_name", "") or
                        getattr(getattr(pair, "product_tmpl_set_id", False), "trade_name", "") or ""
                    )
                    agg[key] = {
                        "num_pedido": order.name,
                        "marca":      self._shoes_get(order, "shoes_campaign_id.product_brand_id.name"),
                        "temporada":  self._shoes_get(order, "shoes_campaign_id.name"),
                        "nombre_web": trade_name,
                        "modelo":     pair.name or "",
                        "color":      self._shoes_get(pair, "color_value_id.name"),
                        "talla":      self._shoes_get(pair, "size_value_id.name"),
                        "quantity":   0,
                        "tip":        self._shoes_get(pair, "categ_id.name"),
                        "barcode":    pair.barcode or "",
                        "desc_corta": self._shoes_get(pair, "shoes_last_id.description"),
                        "mat_ext": self._shoes_mat(
                            task,
                            "shoes_material_external1_id", "shoes_material_external1_percent",
                            "shoes_material_external2_id", "shoes_material_external2_percent",
                        ),
                        "plantilla": self._shoes_mat(
                            task,
                            "shoes_material_lin_internal1_id", "shoes_material_lin_internal1_percent",
                            "shoes_material_lin_internal2_id", "shoes_material_lin_internal2_percent",
                        ),
                        "forro":     self._shoes_get(pair, "shoes_last_id.insole_material_id.name"),
                        "punta":     self._shoes_get(pair, "shoes_last_id.toe"),
                        "tacon":     self._shoes_get(pair, "shoes_last_id.heel_id.name"),
                        "alt_tacon": self._shoes_get(pair, "shoes_last_id.heel_height") or 0,
                        "coste":     line.pair_price,
                        "moneda":    self._shoes_get(order, "currency_id.name"),
                        "origen":    self._shoes_get(pair, "intrastat_duty_id.country_id.code"),
                        "arancel":   self._shoes_get(pair, "intrastat_duty_id.intrastat_id.code"),
                    }
                agg[key]["quantity"] += qty

        return [agg[k] for k in seq]

    # -------------------------------------------------------------------------
    # Hoja Surtidos
    # -------------------------------------------------------------------------

    _SURTIDOS_IMG_COL = next(
        (i + 1 for i, (k, *_) in enumerate(_SURTIDOS_COLS) if k == "imagen"), None
    )
    _SURTIDOS_IMG_PX = 60   # tamaño en píxeles
    _SURTIDOS_ROW_H = 46    # altura de fila en puntos (~60 px)

    def _shoes_write_surtidos(self, wb, lines):
        ws = wb.create_sheet("Surtidos")
        ws.freeze_panes = "A2"
        self._shoes_header(ws, _SURTIDOS_COLS)

        for row_idx, row in enumerate(self._shoes_surtidos_rows(lines), start=2):
            ws.row_dimensions[row_idx].height = self._SURTIDOS_ROW_H
            fill = _FILL_ALT if row_idx % 2 == 0 else None
            for col, (key, _, _, fmt) in enumerate(_SURTIDOS_COLS, start=1):
                if fmt == "image":
                    continue
                val = self._shoes_fmt(row.get(key, ""), fmt)
                cell = ws.cell(row=row_idx, column=col, value=val)
                if fill:
                    cell.fill = fill
                cell.number_format = self._shoes_numfmt(fmt, val)
                cell.alignment = _ALIGN_RIGHT if fmt in ("currency", "number", "integer") else _ALIGN_LEFT

            img_b64 = row.get("_image")
            if img_b64 and XlImage and self._SURTIDOS_IMG_COL:
                try:
                    xl_img = XlImage(io.BytesIO(base64.b64decode(img_b64)))
                    xl_img.width = self._SURTIDOS_IMG_PX
                    xl_img.height = self._SURTIDOS_IMG_PX
                    ws.add_image(xl_img, f"{get_column_letter(self._SURTIDOS_IMG_COL)}{row_idx}")
                except Exception:
                    pass

    def _shoes_surtidos_rows(self, lines):
        rows = []
        for line in lines:
            product = line.product_id
            if not product.is_assortment and not product.is_pair:
                continue
            order = line.order_id

            # Desglose de tallas desde la BoM
            tallas = ""
            if product.is_assortment:
                bom = product.bom_ids[:1]
                if bom:
                    parts = []
                    for bl in bom.bom_line_ids:
                        if bl.product_id.is_pair and bl.product_id.size_value_id:
                            parts.append(
                                f"{bl.product_id.size_value_id.name}x{int(bl.product_qty)}"
                            )
                    tallas = ", ".join(parts)

            assortment_name = self._shoes_get(product, "assortment_attribute_id.name")
            style_code = (product.name or "") + (f" {assortment_name}" if assortment_name else "")

            pares_bul = product.pairs_count or 1

            rows.append({
                "_image":     line.product_id.image_1920 or False,
                "navima_po":  order.name,
                "temporada":  self._shoes_get(order, "shoes_campaign_id.name"),
                "brand":      self._shoes_get(order, "shoes_campaign_id.product_brand_id.name"),
                "modelo":     product.name or "",
                "surtido":    assortment_name,
                "style_code": style_code,
                "color":      self._shoes_get(product, "color_value_id.name"),
                "categoria":  self._shoes_get(product, "categ_id.name"),
                "price":      line.pair_price,
                "moneda":     self._shoes_get(order, "currency_id.name"),
                "tallas":     tallas,
                "par_nota":   line.pairs_count or 0,
                "bultos":     int(line.product_uom_qty),
                "pares_bul":  pares_bul,
                "portes":     self._shoes_get(order, "incoterm.name"),
                "su_ref":     getattr(line, "su_referencia", "") or "",
            })
        return rows

    # -------------------------------------------------------------------------
    # Helpers
    # -------------------------------------------------------------------------

    def _shoes_task(self, product):
        """shoes_task_id del par o, si está vacío, del surtido padre."""
        task = getattr(product, "shoes_task_id", False)
        if not task:
            parent = getattr(product, "product_tmpl_set_id", False)
            if parent:
                task = getattr(parent, "shoes_task_id", False)
        return task or False

    def _shoes_mat(self, task, f1, pct1, f2, pct2):
        """Combina material + porcentaje de dos campos de tarea."""
        if not task:
            return ""
        parts = []
        for fname, pname in ((f1, pct1), (f2, pct2)):
            mat = getattr(task, fname, False)
            if mat:
                pct = getattr(task, pname, 0) or 0
                parts.append(f"{mat.name} {pct:.0f}%" if pct else mat.name)
        return " / ".join(parts)

    def _shoes_get(self, record, path):
        """Resuelve un campo en notación dot-path; devuelve '' si vacío."""
        val = record
        for part in path.split("."):
            if not val or val is False:
                return ""
            if hasattr(val, "ids") and not val.ids:
                return ""
            val = getattr(val, part, "")
        if val is False or val is None:
            return ""
        if hasattr(val, "ids"):
            return ""
        return val

    def _shoes_header(self, ws, cols):
        ws.row_dimensions[1].height = 30
        for col, (_, label, width, _) in enumerate(cols, start=1):
            cell = ws.cell(row=1, column=col, value=label)
            cell.font = _FONT_WHITE_BOLD
            cell.fill = _FILL_HEADER
            cell.alignment = _ALIGN_CENTER
            ws.column_dimensions[get_column_letter(col)].width = width

    def _shoes_fmt(self, val, fmt):
        from datetime import datetime, date as _date
        if isinstance(val, datetime):
            return val.replace(tzinfo=None)
        if isinstance(val, _date):
            return val
        if fmt in ("number", "integer", "currency"):
            if isinstance(val, (int, float)) and not isinstance(val, bool):
                return val
            try:
                return float(val)
            except (TypeError, ValueError):
                return 0
        return str(val) if val != "" else ""

    def _shoes_numfmt(self, fmt, val):
        from datetime import datetime, date as _date
        if fmt == "date" and isinstance(val, (_date, datetime)):
            return "DD/MM/YYYY"
        if fmt in ("currency", "number"):
            return "#,##0.00"
        if fmt == "integer":
            return "#,##0"
        return "General"
